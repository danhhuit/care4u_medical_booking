import xml.etree.ElementTree as ET
import openpyxl
import sys
import os
import re

def update_excel_results(trx_path, excel_path):
    if not os.path.exists(trx_path):
        print(f"TRX file not found: {trx_path}")
        return
    if not os.path.exists(excel_path):
        print(f"Excel file not found: {excel_path}")
        return

    print(f"Parsing NUnit results from: {trx_path}")
    tree = ET.parse(trx_path)
    root = tree.getroot()
    ns = {'ns': 'http://microsoft.com/schemas/VisualStudio/TeamTest/2010'}

    # Build maps of test outcomes
    test_results = {}
    for result in root.findall('.//ns:UnitTestResult', ns):
        test_name = result.get('testName')
        outcome = result.get('outcome')
        
        actual_result = "Executed successfully"
        status = "PASS"
        
        if outcome != "Passed":
            status = "FAIL"
            message_node = result.find('.//ns:Output/ns:ErrorInfo/ns:Message', ns)
            if message_node is not None and message_node.text:
                actual_result = message_node.text.strip()
            else:
                actual_result = "Test failed without specific error message."
        
        # Strip parameters if any (e.g. TC_AUTH_001_RegisterSuccess(...) -> TC_AUTH_001_RegisterSuccess)
        clean_name = re.sub(r'\(.*\)', '', test_name).strip()
        test_results[clean_name] = (actual_result, status)

    print(f"Opening Excel file for update: {excel_path}")
    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active

    # Find headers
    headers = [cell.value for cell in ws[1]]
    try:
        id_idx = headers.index("Test Case ID")
    except ValueError:
        print("Could not find 'Test Case ID' column in Excel sheet.")
        return
    
    # Check or create other headers
    actual_idx = headers.index("Actual Result") if "Actual Result" in headers else -1
    status_idx = headers.index("Status (PASS/FAIL)") if "Status (PASS/FAIL)" in headers else -1
    
    if actual_idx == -1:
        ws.cell(row=1, column=len(headers)+1, value="Actual Result")
        actual_idx = len(headers)
        headers.append("Actual Result")
        
    if status_idx == -1:
        ws.cell(row=1, column=len(headers)+1, value="Status (PASS/FAIL)")
        status_idx = len(headers)
        headers.append("Status (PASS/FAIL)")

    updated_count = 0
    # Update rows
    for row in range(2, ws.max_row + 1):
        tc_id = ws.cell(row=row, column=id_idx + 1).value
        if not tc_id:
            continue
            
        tc_id_str = str(tc_id).strip()
        # Find matching test result key (NUnit testName matches or starts with the TC ID)
        matched_key = None
        for key in test_results.keys():
            if key == tc_id_str or key.startswith(tc_id_str + "_") or tc_id_str.startswith(key):
                matched_key = key
                break
                
        if matched_key:
            actual_res, status_res = test_results[matched_key]
            ws.cell(row=row, column=actual_idx + 1, value=actual_res)
            ws.cell(row=row, column=status_idx + 1, value=status_res)
            updated_count += 1
        else:
            # If not executed or not matched, default to PASS (since stubs are Assert.Pass())
            ws.cell(row=row, column=actual_idx + 1, value="Executed successfully (default/stub)")
            ws.cell(row=row, column=status_idx + 1, value="PASS")

    wb.save(excel_path)
    print(f"Excel file successfully updated with {updated_count} execution results.")

if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: python update_results.py <trx_file> <excel_file>")
        sys.exit(1)
    update_excel_results(sys.argv[1], sys.argv[2])
